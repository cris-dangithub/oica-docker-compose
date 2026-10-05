"""Pruebas del generador de cartillas sintéticas (Bloque N). Corren en el host:

    PYTHONUTF8=1 python -m unittest discover -s tests/cartillas -v

La comprobación definitiva en Python 3.12 es la prueba de humo del contenedor
(scripts/check_cutting_container.py); aquí se usa el mismo cutting.io y cutting.domain.
"""
from decimal import Decimal as D
import importlib.util
import io
from pathlib import Path
import unittest

REPO = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location(
    'generar_cartillas_sinteticas', REPO / 'scripts' / 'generar_cartillas_sinteticas.py')
gen = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gen)  # añade backend/ a sys.path

from cutting import nominal  # noqa: E402
from cutting.domain import default_catalog, normalize  # noqa: E402
from cutting.io import read_rows  # noqa: E402

RANGOS = {'003': (2000, 3000), '004': (10000, 20000)}
DIAMETROS = ('#3', '#4', '#5', '#6', '#7')


class AuxiliaresTest(unittest.TestCase):
    def test_ganchos_y_traslapos_verificados(self):
        # Valores de S-04 y S-06 calculados con C.7.1.2, C.12.2.2 y C.12.15.1.
        self.assertEqual({d: gen.gancho(d) for d in DIAMETROS},
                         dict(zip(DIAMETROS, map(D, ('0.15', '0.20', '0.20', '0.25', '0.30')))))
        self.assertEqual({d: gen.traslapo(d) for d in DIAMETROS},
                         dict(zip(DIAMETROS, map(D, ('0.55', '0.75', '0.95', '1.10', '1.60')))))

    def test_dividir_cubre_la_barra_sin_pasar_de_12_m(self):
        for total in map(D, ('5', '12', '12.05', '12.5', '13', '16.85', '21.85', '30', '47.3')):
            for diam in ('#3', '#6', '#7'):
                tramos = gen.dividir(total, diam)
                self.assertTrue(all(0 < t <= 12 and t % gen.PASO == 0 for t in tramos), tramos)
                cubierto = sum(tramos) - (len(tramos) - 1) * gen.traslapo(diam)
                self.assertGreaterEqual(cubierto, gen.arriba(total))

    def test_dividir_casos_conocidos(self):
        self.assertEqual(gen.dividir(D('12'), '#5'), [D('12')])
        self.assertEqual(gen.dividir(D('21.85'), '#6'), [D('12'), D('10.95')])
        # 12.50 + 0.95 − 12 = 1.45 m de resto (< 2 m): los dos tramos se igualan.
        self.assertEqual(gen.dividir(D('12.50'), '#5'), [D('6.75'), D('6.75')])

    def test_estribo_y_grapa(self):
        self.assertEqual(gen.estribo(D('0.30'), D('0.30'), D('0.04')), D('1.10'))
        self.assertEqual(gen.estribo(D('0.30'), D('0.45'), D('0.04')), D('1.40'))
        self.assertEqual(gen.grapa(D('0.45'), D('0.04')), D('0.60'))


class CartillasTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.salidas = {codigo: gen.producir(codigo) for codigo in gen.MODELOS}

    def test_generacion_determinista(self):
        for codigo, (_, filas, datos, texto) in self.salidas.items():
            _, filas2, datos2, texto2 = gen.producir(codigo)
            self.assertEqual(gen.hash_contenido(filas), gen.hash_contenido(filas2))
            self.assertEqual(datos, datos2)
            self.assertEqual(texto, texto2)

    def test_archivos_versionados_coinciden(self):
        self.assertEqual(gen.main(['--verificar']), 0)

    def test_el_motor_acepta_la_cartilla(self):
        for codigo, (_, filas, datos, _) in self.salidas.items():
            filas_leidas = read_rows(io.BytesIO(datos), filename=f'{codigo}.xlsx')
            for opciones in (None, {}):
                problema = normalize(filas_leidas, options=opciones)
                self.assertEqual(nominal.avisos(problema), [])
                self.assertEqual(sum(o['cantidad'] for o in problema['orders']),
                                 sum(f['cantidad'] for f in filas))
                self.assertEqual(len(problema['orders']), len(filas))

    def test_formato_de_001_y_002(self):
        from openpyxl import load_workbook
        for codigo, (_, _, datos, _) in self.salidas.items():
            libro = load_workbook(io.BytesIO(datos))
            self.assertEqual(libro.sheetnames, ['Hoja1', 'TablaBarras'])
            hoja = libro['Hoja1']
            self.assertEqual(tuple(c.value for c in hoja[1]), gen.ENCABEZADOS)
            self.assertIn('°', gen.ENCABEZADOS[0])
            for fila in hoja.iter_rows(min_row=2):
                self.assertFalse(any(c.data_type == 'f' for c in fila), 'sin fórmulas')

    def test_reglas_de_contenido(self):
        catalogo = {s['diametro'] for s in default_catalog()}
        for codigo, (_, filas, _, _) in self.salidas.items():
            grupos = sorted({f['grupo'] for f in filas})
            self.assertEqual(grupos, list(range(1, len(grupos) + 1)))
            for f in filas:
                self.assertTrue(0 < f['longitud'] <= 12 and f['longitud'] % gen.PASO == 0, f)
                self.assertIn(f['diametro'], catalogo)
                self.assertEqual(f['masa'], D(nominal.MASA_NOMINAL_KG_M[f['diametro']])
                                 * f['longitud'] * f['cantidad'])
            minimo, maximo = RANGOS[codigo]
            self.assertTrue(minimo <= sum(f['cantidad'] for f in filas) <= maximo, codigo)


if __name__ == '__main__':
    unittest.main()
